import json
import glob
import os
import argparse
from typing import Dict, Any, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ----------------------------
# Configuration constants
# ----------------------------
ACTION_CATEGORY_MAP = {
    "Quick": "attack",
    "Heavy": "attack",
    "Special": "attack",
    "ChargeStart": "attack",
    "ChargeRelease": "attack",

    "Dash": "mobility",
    "BackDash": "mobility",
    "Jump": "mobility",
    "MoveLeft": "mobility",
    "MoveRight": "mobility",
    "MoveStop": "mobility",
    "DropPlatform": "mobility",

    "BlockStart": "defense",
    "BlockHold": "defense",
    "Parry": "defense",
    "ParryAttempt": "defense",
    "Dodge": "defense",
    "Roll": "defense",
}

ACTION_SUBTYPE_MAP = {
    "Quick": "quick",
    "Heavy": "heavy",
    "Special": "special",
    "ChargeRelease": "charge",
    "Projectile": "projectile",
}

MOVE_SUBTYPE_MAP = {
    "Quick": "quick",
    "Heavy": "heavy",
    "Special": "special",
    "Charge": "charge",
    "Projectile": "projectile",
}

EMA_ALPHA = 0.35

FEATURE_COLS = [
    "attack_rate",
    "mobility_rate",
    "defense_rate",
    "hit_rate",
    "miss_rate",
    "dps_dealt",
    "dps_taken",
    "block_rate",
    "dodge_rate",
    "quick_rate",
    "heavy_rate",
    "special_rate",
    "charge_rate",
    "projectile_rate",
    "risk_index",
    "parry_rate",
]


# ----------------------------
# Helpers / parsing
# ----------------------------
def parse_args():
    parser = argparse.ArgumentParser(description="Telemetry extractor")
    parser.add_argument("--telemetry-dir", default="Telemetry")
    parser.add_argument("--out-dir", default="out")
    parser.add_argument(
        "--skip-human-xlsx",
        action="store_true",
        help="If set, do not generate the human_readable_export.xlsx file",
    )
    return parser.parse_args()


def safe_get(d: Dict[str, Any], *keys, default=None):
    cur = d
    for k in keys:
        if not isinstance(cur, dict) or k not in cur:
            return default
        cur = cur[k]
    return cur


def infer_seat_ids(meta: Dict[str, Any]) -> Tuple[Optional[str], Optional[str]]:
    p1 = meta.get("p1Id") or meta.get("player1Id") or meta.get("p1") or "P1"
    p2 = meta.get("p2Id") or meta.get("player2Id") or meta.get("p2") or "P2"
    return p1, p2


def infer_profile_info(meta: Dict[str, Any]) -> Tuple[str, str, str, str]:
    p1_profile_id = str(meta.get("p1ProfileId") or "").strip()
    p1_profile_name = str(meta.get("p1ProfileName") or "").strip()

    p2_profile_id = str(meta.get("p2ProfileId") or "").strip()
    p2_profile_name = str(meta.get("p2ProfileName") or "").strip()

    return p1_profile_id, p1_profile_name, p2_profile_id, p2_profile_name


def infer_winner_id(meta: Dict[str, Any]) -> Optional[str]:
    return meta.get("winnerId") or meta.get("winner") or meta.get("winnerPlayerId") or None


def as_float(x, default=0.0):
    try:
        return float(x)
    except Exception:
        return default


def normalize_rate(count: float, duration: float) -> float:
    if duration <= 1e-9:
        return 0.0
    return count / duration


def as_optional_float(x) -> Optional[float]:
    try:
        if x is None or x == "":
            return None
        return float(x)
    except Exception:
        return None


def build_effective_damage_map(events: List[Dict[str, Any]]) -> Dict[int, float]:
    groups: Dict[Tuple[Any, ...], List[Tuple[int, Dict[str, Any], float]]] = {}

    for idx, e in enumerate(events):
        if not isinstance(e, dict):
            continue

        etype = e.get("eventType") or e.get("type")
        if etype != "DamageApplied":
            continue

        raw_damage = max(0.0, as_float(e.get("actualDamage") or e.get("finalDamage") or 0.0))
        defender = e.get("defenderId") or e.get("defender") or ""
        frame = e.get("frame")
        t = as_optional_float(e.get("t"))

        if defender and frame not in (None, ""):
            key = ("frame", defender, str(frame))
        elif defender and t is not None:
            key = ("time", defender, round(t, 3))
        else:
            key = ("event", idx)

        groups.setdefault(key, []).append((idx, e, raw_damage))

    effective_damage_by_index: Dict[int, float] = {}

    for entries in groups.values():
        valid_hp_entries = [
            (idx, e, raw_damage)
            for idx, e, raw_damage in entries
            if as_optional_float(e.get("hpDefenderBefore")) is not None
            and as_optional_float(e.get("hpDefenderAfter")) is not None
        ]

        if not valid_hp_entries:
            for idx, _e, raw_damage in entries:
                effective_damage_by_index[idx] = raw_damage
            continue

        group_before = max(
            max(0.0, as_float(e.get("hpDefenderBefore")))
            for _idx, e, _raw_damage in valid_hp_entries
        )
        group_after = min(
            max(0.0, as_float(e.get("hpDefenderAfter")))
            for _idx, e, _raw_damage in valid_hp_entries
        )
        group_effective_damage = max(0.0, group_before - group_after)
        group_raw_damage = sum(raw_damage for _idx, _e, raw_damage in entries)

        if group_raw_damage <= 1e-9:
            for idx, _e, _raw_damage in entries:
                effective_damage_by_index[idx] = 0.0
            continue

        for idx, _e, raw_damage in entries:
            effective_damage_by_index[idx] = group_effective_damage * (raw_damage / group_raw_damage)

    return effective_damage_by_index


def ratio(num: float, den: float) -> float:
    if den <= 1e-9:
        return 0.0
    return num / den


def export_human_readable_xlsx(round_df: pd.DataFrame, profile_df: pd.DataFrame, xlsx_path: str) -> None:
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        round_df.to_excel(writer, sheet_name="round_level", index=False)
        profile_df.to_excel(writer, sheet_name="profile_level", index=False)

    wb = load_workbook(xlsx_path)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    rate_fill = PatternFill(fill_type="solid", fgColor="EAF7EA")
    id_fill = PatternFill(fill_type="solid", fgColor="F4EAEA")
    metric_fill = PatternFill(fill_type="solid", fgColor="FFF4CC")

    def style_sheet(ws):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for col_idx, column_cells in enumerate(ws.columns, start=1):
            col_letter = get_column_letter(col_idx)
            header = ws.cell(row=1, column=col_idx).value
            header_str = str(header).strip() if header is not None else ""

            max_length = len(header_str)
            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))

            ws.column_dimensions[col_letter].width = min(max_length + 2, 28)

            if header_str.endswith("_id") or header_str in {
                "match_id", "profile_id", "player_id", "slot_id"
            }:
                for cell in column_cells:
                    cell.fill = id_fill

            elif "rate" in header_str or header_str in {
                "hit_rate", "miss_rate", "block_rate", "dodge_rate"
            }:
                for cell in column_cells:
                    cell.fill = rate_fill
                    if cell.row > 1 and isinstance(cell.value, (int, float)):
                        cell.number_format = "0.0000"

            elif header_str in {
                "damage_dealt", "damage_taken", "dps_dealt", "dps_taken",
                "actions_total", "actions_attack", "actions_mobility",
                "actions_defense", "actions_quick", "actions_heavy",
                "actions_special", "actions_charge", "actions_projectile",
                "hit_attempts", "misses", "hits_landed", "hits_received",
                "blocked_hits_as_defender", "dodged_hits_as_defender",
                "risk_index"
            }:
                for cell in column_cells:
                    cell.fill = metric_fill
                    if cell.row > 1 and isinstance(cell.value, float):
                        cell.number_format = "0.0000"

        ws.sheet_view.showGridLines = True

    for sheet_name in wb.sheetnames:
        style_sheet(wb[sheet_name])

    wb.save(xlsx_path)


# ----------------------------
# Core feature extraction per round
# ----------------------------
def extract_round_features(doc: Dict[str, Any], filename: str) -> List[Dict[str, Any]]:
    meta = doc.get("meta", {}) if isinstance(doc.get("meta", {}), dict) else {}
    events = doc.get("events", []) if isinstance(doc.get("events", []), list) else []

    duration = as_float(doc.get("durationSeconds") or meta.get("durationSeconds") or 0.0)
    match_id = doc.get("matchId") or meta.get("matchId") or os.path.basename(filename)

    round_number = meta.get("roundNumber", None)
    stage = meta.get("map") or meta.get("stage") or meta.get("mapName") or None
    mode = meta.get("mode") or None

    p1_seat, p2_seat = infer_seat_ids(meta)
    winner_id = infer_winner_id(meta)

    p1_profile_id, p1_profile_name, p2_profile_id, p2_profile_name = infer_profile_info(meta)

    def identity_for_slot(slot_id: str):
        if slot_id == p1_seat:
            return p1_profile_id, p1_profile_name
        if slot_id == p2_seat:
            return p2_profile_id, p2_profile_name
        return "", ""

    def player_key(slot_id: str, profile_id: str):
        if profile_id and profile_id not in ("GUEST", "NONE", "UNKNOWN"):
            return profile_id
        return slot_id

    agg: Dict[str, Dict[str, float]] = {}

    def ensure(pid: str):
        if pid not in agg:
            agg[pid] = {
                "actions_total": 0,
                "actions_attack": 0,
                "actions_mobility": 0,
                "actions_defense": 0,
                "actions_quick": 0,
                "actions_heavy": 0,
                "actions_special": 0,
                "actions_charge": 0,
                "actions_projectile": 0,
                "actions_parry": 0,
                "hit_attempts": 0,
                "misses": 0,
                "damage_dealt": 0.0,
                "damage_taken": 0.0,
                "blocked_hits_as_defender": 0,
                "dodged_hits_as_defender": 0,
                "hits_received": 0,
                "hits_landed": 0,
            }

    effective_damage_by_event = build_effective_damage_map(events)

    for event_index, e in enumerate(events):
        etype = e.get("eventType") or e.get("type")

        if etype == "Action":
            pid = e.get("playerId") or e.get("actorId") or e.get("attackerId")
            if not pid:
                continue

            ensure(pid)
            agg[pid]["actions_total"] += 1

            action_type = e.get("actionType") or e.get("action") or ""

            if action_type == "Parry":
                agg[pid]["actions_parry"] += 1
            cat = ACTION_CATEGORY_MAP.get(action_type, None)
            if cat == "attack":
                agg[pid]["actions_attack"] += 1
            elif cat == "mobility":
                agg[pid]["actions_mobility"] += 1
            elif cat == "defense":
                agg[pid]["actions_defense"] += 1

            subtype = ACTION_SUBTYPE_MAP.get(action_type, None)
            if subtype == "quick":
                agg[pid]["actions_quick"] += 1
            elif subtype == "heavy":
                agg[pid]["actions_heavy"] += 1
            elif subtype == "special":
                agg[pid]["actions_special"] += 1
            elif subtype == "charge":
                agg[pid]["actions_charge"] += 1
            elif subtype == "projectile":
                agg[pid]["actions_projectile"] += 1

        elif etype == "HitAttempt":
            attacker = e.get("attackerId") or e.get("playerId")
            if not attacker:
                continue

            ensure(attacker)
            agg[attacker]["hit_attempts"] += 1

            _move = e.get("moveType") or ""
            _ = MOVE_SUBTYPE_MAP.get(_move, None)

        elif etype == "Miss":
            attacker = e.get("attackerId") or e.get("playerId")
            if not attacker:
                continue

            ensure(attacker)
            agg[attacker]["misses"] += 1

        elif etype == "DamageApplied":
            attacker = e.get("attackerId")
            defender = e.get("defenderId")
            raw_dmg = as_float(e.get("actualDamage") or e.get("finalDamage") or 0.0)
            dmg = effective_damage_by_event.get(event_index, max(0.0, raw_dmg))
            blocked = bool(e.get("blocked") or e.get("wasBlocked") or False)
            dodged = bool(e.get("dodged") or e.get("wasDodged") or False)

            move_type = str(e.get("moveType", "")).strip()
            source_type = str(e.get("sourceType", "")).strip()
            is_dot = (move_type == "PoisonTick") or (source_type == "Dot")

            if attacker:
                ensure(attacker)
                agg[attacker]["damage_dealt"] += dmg
                if raw_dmg > 0 and not is_dot:
                    agg[attacker]["hits_landed"] += 1

            if defender:
                ensure(defender)
                agg[defender]["damage_taken"] += dmg
                if not is_dot:
                    agg[defender]["hits_received"] += 1
                    if blocked:
                        agg[defender]["blocked_hits_as_defender"] += 1
                    if dodged:
                        agg[defender]["dodged_hits_as_defender"] += 1

    for pid in [p1_seat, p2_seat]:
        if pid:
            ensure(pid)

    rows: List[Dict[str, Any]] = []
    MIN_ACTIONS = 25
    MIN_DURATION = 5
    INVALID_PROFILES = {"GUEST", "UNKNOWN", "", "NONE", None}

    if (p1_profile_id in {"GUEST", "UNKNOWN", "NONE", "", None} and
        p2_profile_id in {"GUEST", "UNKNOWN", "NONE", "", None}):
        return []

    for pid, a in agg.items():
        profile_id, profile_name = identity_for_slot(pid)

        if profile_id in INVALID_PROFILES:
            continue

        if a["actions_total"] < MIN_ACTIONS or duration < MIN_DURATION:
            continue

        attack_rate = ratio(a["actions_attack"], a["actions_total"])
        mobility_rate = ratio(a["actions_mobility"], a["actions_total"])
        defense_rate = ratio(a["actions_defense"], a["actions_total"])

        quick_rate = ratio(a["actions_quick"], a["actions_total"])
        heavy_rate = ratio(a["actions_heavy"], a["actions_total"])
        special_rate = ratio(a["actions_special"], a["actions_total"])
        charge_rate = ratio(a["actions_charge"], a["actions_total"])
        projectile_rate = ratio(a["actions_projectile"], a["actions_total"])
        parry_rate = ratio(a["actions_parry"], a["actions_total"])
        total_offensive_outcomes = a["hits_landed"] + a["misses"]
        hit_rate = ratio(a["hits_landed"], total_offensive_outcomes)
        miss_rate = ratio(a["misses"], total_offensive_outcomes)

        dps_dealt = normalize_rate(a["damage_dealt"], duration)
        dps_taken = normalize_rate(a["damage_taken"], duration)

        block_rate = ratio(a["blocked_hits_as_defender"], a["hits_received"])
        dodge_rate = ratio(a["dodged_hits_as_defender"], a["hits_received"])

        risk_index = (
            0.20 * miss_rate +
            0.40 * parry_rate +
            0.40 * charge_rate
        )

        rows.append({
            "match_id": match_id,
            "round_number": round_number,
            "stage": stage,
            "mode": mode,
            "duration_s": duration,
            "slot_id": pid,
            "is_p1": (pid == p1_seat) if p1_seat else None,
            "is_p2": (pid == p2_seat) if p2_seat else None,
            "won_round": (pid == winner_id) if winner_id else None,
            "profile_id": profile_id,
            "profile_name": profile_name,
            "player_id": player_key(pid, profile_id),
            "actions_total": a["actions_total"],
            "actions_attack": a["actions_attack"],
            "actions_mobility": a["actions_mobility"],
            "actions_defense": a["actions_defense"],
            "actions_quick": a["actions_quick"],
            "actions_heavy": a["actions_heavy"],
            "actions_special": a["actions_special"],
            "actions_charge": a["actions_charge"],
            "actions_projectile": a["actions_projectile"],
            "hit_attempts": a["hit_attempts"],
            "misses": a["misses"],
            "hits_landed": a["hits_landed"],
            "hits_received": a["hits_received"],
            "damage_dealt": a["damage_dealt"],
            "damage_taken": a["damage_taken"],
            "blocked_hits_as_defender": a["blocked_hits_as_defender"],
            "dodged_hits_as_defender": a["dodged_hits_as_defender"],
            "attack_rate": attack_rate,
            "mobility_rate": mobility_rate,
            "defense_rate": defense_rate,
            "hit_rate": hit_rate,
            "miss_rate": miss_rate,
            "dps_dealt": dps_dealt,
            "dps_taken": dps_taken,
            "block_rate": block_rate,
            "dodge_rate": dodge_rate,
            "quick_rate": quick_rate,
            "heavy_rate": heavy_rate,
            "special_rate": special_rate,
            "charge_rate": charge_rate,
            "projectile_rate": projectile_rate,
            "risk_index": risk_index,
            "parry_rate": parry_rate,
        })

    return rows


# ----------------------------
# Profiling layer (EMA per player)
# ----------------------------
def build_ema_profiles(df_round: pd.DataFrame, alpha: float = EMA_ALPHA) -> pd.DataFrame:
    sort_cols = ["player_id"]
    if "round_number" in df_round.columns and df_round["round_number"].notna().any():
        sort_cols.append("round_number")
    else:
        sort_cols.append("match_id")

    df = df_round.sort_values(sort_cols).copy()

    out_rows = []
    for pid, g in df.groupby("player_id", sort=False):
        g = g.copy()
        ema = None

        for _, row in g.iterrows():
            x = row[FEATURE_COLS].astype(float).to_numpy()
            if ema is None:
                ema = x
            else:
                ema = alpha * x + (1 - alpha) * ema

            r = row.to_dict()
            for i, col in enumerate(FEATURE_COLS):
                r[f"ema_{col}"] = float(ema[i])
            out_rows.append(r)

    return pd.DataFrame(out_rows)


def run_extractor(
    telemetry_dir: str,
    out_dir: str,
    skip_human_xlsx: bool = False,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    round_csv = os.path.join(out_dir, "dataset_round_level.csv")
    profile_csv = os.path.join(out_dir, "dataset_profile_level.csv")
    human_xlsx = os.path.join(out_dir, "human_readable_export.xlsx")

    os.makedirs(out_dir, exist_ok=True)

    files = sorted(glob.glob(os.path.join(telemetry_dir, "*.json")))
    if not files:
        raise FileNotFoundError(f"No JSON files found in: {telemetry_dir}")

    all_rows = []
    bad_files = 0

    for fp in files:
        try:
            with open(fp, "r", encoding="utf-8") as f:
                doc = json.load(f)
            rows = extract_round_features(doc, fp)
            all_rows.extend(rows)
        except Exception as ex:
            bad_files += 1
            print(f"[WARN] Failed parsing {fp}: {ex}")

    if not all_rows:
        raise RuntimeError("Parsed 0 rows. Check schema / Telemetry JSON content.")

    df_round = pd.DataFrame(all_rows)

    if "won_round" in df_round.columns:
        df_round["won_round"] = df_round["won_round"].astype("float")

    df_round.to_csv(round_csv, index=False)
    print(f"[OK] Round-level dataset saved: {round_csv}  (rows={len(df_round)})")

    if bad_files:
        print(f"[INFO] Bad files skipped: {bad_files}")

    MIN_MATCHES_PER_PROFILE = 5
    profile_match_counts = df_round.groupby("profile_id").size().reset_index(name="match_count")

    print("\n--- profile_match_counts ---")
    print(profile_match_counts.sort_values("match_count", ascending=False))

    valid_profiles = set(
        profile_match_counts.loc[
            profile_match_counts["match_count"] >= MIN_MATCHES_PER_PROFILE,
            "profile_id"
        ]
    )

    df_profile_source = df_round[df_round["profile_id"].isin(valid_profiles)].copy()

    print(f"[DEBUG] valid_profiles={len(valid_profiles)}")
    print(f"[DEBUG] df_profile_source rows={len(df_profile_source)}")

    if df_profile_source.empty:
        print("[WARN] No profiles passed MIN_MATCHES_PER_PROFILE filter. Saving empty profile dataset.")
        df_profile = pd.DataFrame()
    else:
        df_profile = build_ema_profiles(df_profile_source, alpha=EMA_ALPHA)

    df_profile.to_csv(profile_csv, index=False)
    print(f"[OK] Profile-level dataset saved: {profile_csv}  (rows={len(df_profile)})")

    if not skip_human_xlsx:
        export_human_readable_xlsx(df_round, df_profile, human_xlsx)
        print(f"[OK] Human-readable Excel export saved: {human_xlsx}")
    else:
        print("[INFO] Skipping human_readable_export.xlsx")

    print("\n--- Feature means (round-level) ---")
    existing_feature_cols = [c for c in FEATURE_COLS if c in df_round.columns]
    if existing_feature_cols:
        print(df_round[existing_feature_cols].mean(numeric_only=True).round(4))

    print("\n--- Detailed action totals (round-level) ---")
    detail_cols = [
        "actions_quick",
        "actions_heavy",
        "actions_special",
        "actions_charge",
        "actions_projectile",
    ]
    existing_detail_cols = [c for c in detail_cols if c in df_round.columns]
    if existing_detail_cols:
        print(df_round[existing_detail_cols].sum(numeric_only=True))

    return df_round, df_profile


def main_with_args(
    telemetry_dir: str,
    out_dir: str,
    skip_human_xlsx: bool = False,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    return run_extractor(
        telemetry_dir=telemetry_dir,
        out_dir=out_dir,
        skip_human_xlsx=skip_human_xlsx,
    )


# ----------------------------
# Main
# ----------------------------
def main():
    args = parse_args()
    run_extractor(
        telemetry_dir=args.telemetry_dir,
        out_dir=args.out_dir,
        skip_human_xlsx=args.skip_human_xlsx,
    )


if __name__ == "__main__":
    main()
